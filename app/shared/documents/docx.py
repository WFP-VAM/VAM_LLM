"""Word rendering of report blocks. Each drafter gives its look (a WordTheme) and its words (WordLabels).

Without a theme, blocks use python-docx's template: built-in headings and default paragraphs. A drafter's theme
names a paragraph style per role, prepares the document (style definitions, page, header and footer), reads each
block's layout hints and renders the drafter's own table kinds. The Word helpers below are for themes and tables.
"""
from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Mapping, Optional, Sequence

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from .blocks import ReportBlock

HEADING_ROLES = {"title", "major_section", "subsection", "minor_heading"}


def _no_layout(block: ReportBlock) -> Dict[str, Any]:
    return {}


@dataclass(frozen=True)
class WordLabels:
    """The words the renderer writes itself."""

    references: str = "References"
    limitation: str = "Data limitation"
    methodology: str = "Methodology"


@dataclass(frozen=True)
class WordTheme:
    """A drafter's look in Word.

    styles: the paragraph style of each role: the HEADING_ROLES, "body", "caption" and "notice". Without styles,
        headings are python-docx's built-in ones and paragraphs keep the default style.
    setup: prepares the new document before any block is added.
    layout: a block's layout hints (role, page_break_before, keep_with_next, keep_together, compact_after).
    tables: the renderer of each table kind (ReportBlock.meta["table_kind"]); tables of other kinds are left out.
    """

    styles: Mapping[str, str] = field(default_factory=dict)
    setup: Optional[Callable[[Any, Sequence[ReportBlock]], None]] = None
    layout: Callable[[ReportBlock], Dict[str, Any]] = _no_layout
    tables: Mapping[str, Callable[[Any, Dict[str, Any]], None]] = field(default_factory=dict)


def define_paragraph_style(
    doc: Any,
    name: str,
    *,
    font: str,
    size: float,
    color: str = "000000",
    bold: bool = False,
    italic: bool = False,
    before: float = 0.0,
    after: float = 0.0,
    line_spacing: float = 1.0,
    keep_with_next: bool = False,
    keep_together: bool = False,
    outline_level: Optional[int] = None,
) -> Any:
    """Add the named paragraph style, or update it if the document already has it."""
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = font
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    style.font.bold = bold
    style.font.italic = italic
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    fmt.keep_with_next = keep_with_next
    fmt.keep_together = keep_together
    if outline_level is not None:
        properties = style._element.get_or_add_pPr()
        existing = properties.find(qn("w:outlineLvl"))
        if existing is None:
            existing = OxmlElement("w:outlineLvl")
            properties.append(existing)
        existing.set(qn("w:val"), str(int(outline_level)))
    return style


def add_field(paragraph: Any, instruction: str, placeholder: str = "1") -> None:
    """Append a Word field, such as PAGE or NUMPAGES, that Word computes when it opens the file."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def apply_layout(paragraph: Any, layout: Dict[str, Any]) -> None:
    """Apply a block's layout hints to one of its paragraphs; no hints, no change."""
    if not layout:
        return
    paragraph.paragraph_format.page_break_before = bool(
        layout.get("page_break_before")
    )
    paragraph.paragraph_format.keep_with_next = bool(layout.get("keep_with_next"))
    paragraph.paragraph_format.keep_together = bool(layout.get("keep_together"))
    if layout.get("compact_after"):
        paragraph.paragraph_format.space_after = Pt(2)


def shade_cell(cell: Any, rgb_tuple: tuple[int, int, int]) -> None:
    r, g, b = rgb_tuple
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{r:02x}{g:02x}{b:02x}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def repeat_header_row(row: Any) -> None:
    """Repeat the row at the top of every page the table spans."""
    properties = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    properties.append(element)


def keep_row_together(row: Any) -> None:
    """Never split the row across two pages."""
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def set_cell_width(cell: Any, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    properties = cell._tc.get_or_add_tcPr()
    width = properties.get_or_add_tcW()
    width.set(qn("w:w"), str(int(round(width_inches * 1440))))
    width.set(qn("w:type"), "dxa")


def _safe_filename(filename: str) -> str:
    name = filename.strip() or "export.docx"
    name = name.replace("\\", "_").replace("/", "_")
    name = re.sub(r"[^A-Za-z0-9._\- ]+", "_", name)
    if not name.lower().endswith(".docx"):
        name = name + ".docx"
    return name


def build_content_disposition(filename: str) -> str:
    safe = _safe_filename(filename)
    return f'attachment; filename="{safe}"'


def _add_text_lines(
    doc: Any,
    text: str,
    *,
    style: Optional[str] = None,
    layout: Optional[Dict[str, Any]] = None,
) -> None:
    for raw_line in (text or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if line.startswith("- ") or line.startswith("* "):
            paragraph = doc.add_paragraph(line[2:].strip(), style="List Bullet")
            apply_layout(paragraph, layout or {})
            continue

        if re.match(r"^\d+\.\s+", line):
            paragraph = doc.add_paragraph(
                re.sub(r"^\d+\.\s+", "", line), style="List Number"
            )
            apply_layout(paragraph, layout or {})
            continue

        paragraph = doc.add_paragraph(line, style=style)
        apply_layout(paragraph, layout or {})


def _add_notice_box(
    doc: Any,
    text: str,
    *,
    label: str,
    fill: str,
    color: tuple[int, int, int],
    style: Optional[str] = None,
    layout: Optional[Dict[str, Any]] = None,
) -> None:
    cleaned = (text or "").strip()
    if not cleaned:
        return
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    keep_row_together(table.rows[0])
    cell = table.rows[0].cells[0]
    paragraph = cell.paragraphs[0]
    if style:
        paragraph.style = style
        apply_layout(paragraph, layout or {})
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.color.rgb = RGBColor(*color)
    text_run = paragraph.add_run(cleaned)
    text_run.font.size = Pt(9)
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>')
    cell._tc.get_or_add_tcPr().append(shading)
    if not style:
        doc.add_paragraph()


def render_docx(
    report_blocks: Sequence[ReportBlock],
    *,
    theme: WordTheme = WordTheme(),
    labels: WordLabels = WordLabels(),
    visualizations: Optional[Mapping[str, str]] = None,
    include_sources: bool = True,
    include_visualizations: bool = True,
) -> bytes:
    """The Word file of the blocks. Figures are base64 images in visualizations, by figure id."""
    doc = Document()
    visualizations = visualizations or {}
    styles = theme.styles
    if theme.setup is not None:
        theme.setup(doc, report_blocks)

    for block in report_blocks:
        layout = theme.layout(block)
        if block.type == "heading":
            level = int(block.level or 1)
            level = min(max(level, 1), 9)
            if styles and layout:
                role = str(layout.get("role") or "subsection")
                style_role = role if role in HEADING_ROLES else "subsection"
                paragraph = doc.add_paragraph(
                    block.text or "",
                    style=styles[style_role],
                )
                apply_layout(paragraph, layout)
            else:
                doc.add_heading(block.text or "", level=level)
            continue

        if block.type == "paragraph":
            _add_text_lines(
                doc,
                block.text or "",
                style=styles["body"] if styles and layout else None,
                layout=layout,
            )
            continue

        if block.type == "figure":
            if not include_visualizations:
                continue

            fig_id = (block.figure_id or "").strip()
            fig_b64 = visualizations.get(fig_id)
            if not fig_id or not fig_b64:
                continue

            try:
                img_bytes = base64.b64decode(fig_b64)
            except Exception:
                continue

            width = float(block.width) if block.width is not None else 6.0
            width = max(1.0, min(width, 7.0))

            buf = io.BytesIO(img_bytes)
            p = doc.add_paragraph()
            if styles:
                p.style = styles["body"]
            run = p.add_run()
            run.add_picture(buf, width=Inches(width))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            apply_layout(p, layout)

            if block.caption:
                cap = doc.add_paragraph(
                    block.caption,
                    style=styles["caption"] if styles else None,
                )
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if styles:
                    cap.paragraph_format.keep_together = True
            continue

        if block.type == "references":
            if not include_sources:
                continue

            refs = block.references or []
            if not refs:
                continue

            if styles:
                heading = doc.add_paragraph(labels.references, style=styles["subsection"])
                heading.paragraph_format.keep_with_next = True
            else:
                doc.add_heading(labels.references, level=2)
            for ref in refs:
                if not isinstance(ref, dict):
                    continue

                doc_id = str(ref.get("doc_id", "")).strip()
                source = str(ref.get("source", "")).strip()
                date = str(ref.get("date", "")).strip()
                title = str(ref.get("title", "")).strip()
                url = str(ref.get("url", "")).strip()

                parts: List[str] = []
                if doc_id:
                    parts.append(f"[{doc_id}]")
                if source:
                    parts.append(source)
                if date:
                    parts.append(f"({date})")
                if title:
                    parts.append(title)

                paragraph = doc.add_paragraph(
                    " ".join(parts).strip(),
                    style="List Number",
                )
                if styles:
                    paragraph.paragraph_format.space_after = Pt(2)
                if url:
                    doc.add_paragraph(url, style=styles["body"] if styles else None)
            continue

        if block.type == "table":
            meta = block.meta or {}
            kind = meta.get("table_kind") if isinstance(meta, dict) else None
            render_table = theme.tables.get(kind) if isinstance(kind, str) else None
            if render_table is not None:
                render_table(doc, meta)
            continue

        if block.type == "limitation_box":
            _add_notice_box(
                doc,
                block.text or "",
                label=labels.limitation,
                fill="FFF4CC",
                color=(145, 94, 0),
                style=styles["notice"] if styles else None,
                layout=layout,
            )
            continue

        if block.type == "methodology_note":
            _add_notice_box(
                doc,
                block.text or "",
                label=labels.methodology,
                fill="E6F3FF",
                color=(0, 114, 188),
                style=styles["notice"] if styles else None,
                layout=layout,
            )
            continue

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.read()
