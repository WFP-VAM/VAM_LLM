"""The MFI report's layout: the layout contract its blocks carry, and the Word theme that follows it."""
from __future__ import annotations

from typing import Any, Dict, List, Literal, Optional, Sequence

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt, RGBColor
from pydantic import BaseModel, ConfigDict

from app.shared.documents.blocks import ReportBlock
from app.shared.documents.docx import (
    WordTheme,
    add_field,
    define_paragraph_style,
    keep_row_together,
    repeat_header_row,
    set_cell_width,
    shade_cell,
)


def resolve_mfi_report_blocks(result: Dict[str, Any]) -> List[ReportBlock]:
    """Return the report blocks stored with a completed MFI result."""
    stored = result.get("report_blocks")
    if not isinstance(stored, list) or not stored:
        raise ValueError("This MFI result has no report blocks; reports from the previous workflow are no longer supported")
    return [
        item if isinstance(item, ReportBlock) else ReportBlock.model_validate(item)
        for item in stored
    ]


class MFIReportLayoutHint(BaseModel):
    """Typed internal layout contract shared by MFI renderers."""

    model_config = ConfigDict(frozen=True)

    report_family: Literal["mfi"] = "mfi"
    role: Literal[
        "title",
        "major_section",
        "subsection",
        "minor_heading",
        "body",
        "figure",
        "table",
        "notice",
        "references",
    ]
    group_id: Optional[str] = None
    page_break_before: bool = False
    keep_with_next: bool = False
    keep_together: bool = False
    compact_after: bool = False
    country: Optional[str] = None
    methodology_version: Optional[str] = None


_MAJOR_PAGE_BREAK_HEADINGS = {"Executive summary", "MFI dimensions"}


def apply_mfi_layout_contract(
    blocks: List[ReportBlock],
    *,
    country: str,
    methodology_version: str,
) -> List[ReportBlock]:
    """Attach explicit R8 layout roles without making renderers infer semantics."""
    for block in blocks:
        meta = dict(block.meta or {})
        role: str
        page_break_before = False
        keep_with_next = False
        keep_together = False
        compact_after = False
        if block.type == "heading":
            if int(block.level or 1) <= 1:
                role = "title"
            elif int(block.level or 2) == 2:
                role = "major_section"
                page_break_before = str(block.text or "") in (
                    _MAJOR_PAGE_BREAK_HEADINGS
                )
            elif int(block.level or 3) == 3:
                role = "subsection"
            else:
                role = "minor_heading"
            keep_with_next = True
        elif block.type == "figure":
            role = "figure"
            keep_with_next = bool(block.caption)
            keep_together = True
            compact_after = True
        elif block.type == "table":
            role = "table"
            compact_after = True
        elif block.type == "references":
            role = "references"
        elif block.type in {"limitation_box", "methodology_note"}:
            role = "notice"
            keep_together = True
            compact_after = True
        else:
            role = "body"
        layout = MFIReportLayoutHint(
            role=role,
            page_break_before=page_break_before,
            keep_with_next=keep_with_next,
            keep_together=keep_together,
            compact_after=compact_after,
            country=country or None if role == "title" else None,
            methodology_version=(
                methodology_version if role == "title" else None
            ),
        )
        meta["mfi_layout"] = layout.model_dump(mode="json")
        block.meta = meta
    return blocks


def layout_hints(block: ReportBlock) -> Dict[str, Any]:
    """The block's layout contract, as stored in its meta."""
    meta = block.meta or {}
    layout = meta.get("mfi_layout") if isinstance(meta, dict) else None
    return dict(layout) if isinstance(layout, dict) else {}


# --- Word ------------------------------------------------------------------------------------------------------

STYLE_NAMES = {
    "title": "MFI Title",
    "major_section": "MFI Major Heading",
    "subsection": "MFI Subsection Heading",
    "minor_heading": "MFI Minor Heading",
    "body": "MFI Body",
    "notice": "MFI Notice",
    "caption": "MFI Caption",
    "header": "MFI Header",
    "footer": "MFI Footer",
    "table_header": "MFI Table Header",
    "table_body": "MFI Table Body",
}


def _style(doc: Any, role: str, **settings: Any) -> None:
    define_paragraph_style(doc, STYLE_NAMES[role], font="Arial", **settings)


def _define_styles(doc: Any) -> None:
    _style(doc, "title", size=20, color="0072BC", bold=True, after=6, keep_with_next=True, outline_level=0)
    _style(doc, "major_section", size=14, color="0072BC", bold=True, before=12, after=6, keep_with_next=True, outline_level=1)
    _style(doc, "subsection", size=11, color="1F4D78", bold=True, before=8, after=4, keep_with_next=True, outline_level=2)
    _style(doc, "minor_heading", size=9.5, color="1F4D78", bold=True, before=6, after=3, keep_with_next=True, outline_level=3)
    _style(doc, "body", size=9, after=3, line_spacing=1.05)
    _style(doc, "notice", size=8.5, after=2, line_spacing=1.05, keep_together=True)
    _style(doc, "caption", size=8, color="505050", italic=True, after=4, keep_together=True)
    _style(doc, "header", size=8, color="666666", after=0)
    _style(doc, "footer", size=8, color="666666", after=0)
    _style(doc, "table_header", size=7, color="FFFFFF", bold=True, after=0, keep_together=True)
    _style(doc, "table_body", size=7, after=0, line_spacing=1.0, keep_together=True)


def _prepare_document(doc: Any, blocks: Sequence[ReportBlock]) -> None:
    """Styles, A4 page, the header naming the country and the "Page X of Y" footer."""
    title = next(
        (layout_hints(block) for block in blocks if layout_hints(block).get("role") == "title"),
        {},
    )
    _define_styles(doc)
    country = str(title.get("country") or "").strip()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.header_distance = Inches(0.3)
        section.footer_distance = Inches(0.3)
        header = section.header.paragraphs[0]
        header.style = STYLE_NAMES["header"]
        header.text = "MFI Drafter 2.0" + (f" - {country}" if country else "")
        footer = section.footer.paragraphs[0]
        footer.style = STYLE_NAMES["footer"]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Page ")
        add_field(footer, "PAGE")
        footer.add_run(" of ")
        add_field(footer, "NUMPAGES")


def _add_presentation_table(doc: Any, meta: Dict[str, Any]) -> None:
    """Render an already projected MFI table without inferring or formatting cells."""
    rows = meta.get("rows") or []
    if not isinstance(rows, list) or not rows:
        return
    spec_id = str(meta.get("spec_id") or "").strip()
    columns = [str(column) for column in meta.get("columns", []) or []]
    column_specs = meta.get("column_specs") or []
    if not spec_id or not columns or not isinstance(column_specs, list):
        raise ValueError("Projected MFI tables require a spec ID and explicit columns")
    if len(columns) > 8 or len(column_specs) != len(columns):
        raise ValueError("Projected MFI table columns violate the renderer contract")
    spec_keys = [
        str(item.get("key") or "") if isinstance(item, dict) else ""
        for item in column_specs
    ]
    if spec_keys != columns:
        raise ValueError("Projected MFI table columns do not match column_specs")
    values = []
    for row in rows:
        if not isinstance(row, dict) or not isinstance(row.get("values"), dict):
            raise ValueError("Projected MFI table contains a malformed row")
        if any(column not in row["values"] for column in columns):
            raise ValueError("Projected MFI table row is missing a visible column")
        values.append(row["values"])
    title = str(meta.get("title") or "").strip()
    if title:
        title_paragraph = doc.add_paragraph(
            title,
            style=STYLE_NAMES["minor_heading"],
        )
        title_paragraph.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=len(values) + 1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_width = sum(float(item["width_hint"]) for item in column_specs)
    widths = [6.5 * float(item["width_hint"]) / total_width for item in column_specs]
    alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    for index, column in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.text = str(column_specs[index].get("label") or "")
        set_cell_width(cell, widths[index])
        shade_cell(cell, (0, 114, 188))
        for paragraph in cell.paragraphs:
            paragraph.style = STYLE_NAMES["table_header"]
            paragraph.alignment = alignment.get(
                str(column_specs[index].get("alignment") or "left"),
                WD_ALIGN_PARAGRAPH.LEFT,
            )
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(7)
                run.font.color.rgb = RGBColor(255, 255, 255)
    repeat_header_row(table.rows[0])
    for row_index, row in enumerate(values, start=1):
        keep_row_together(table.rows[row_index])
        for column_index, column in enumerate(columns):
            cell = table.rows[row_index].cells[column_index]
            set_cell_width(cell, widths[column_index])
            cell.text = str(row[column])
            for paragraph in cell.paragraphs:
                paragraph.style = STYLE_NAMES["table_body"]
                paragraph.alignment = alignment.get(
                    str(column_specs[column_index].get("alignment") or "left"),
                    WD_ALIGN_PARAGRAPH.LEFT,
                )
                for run in paragraph.runs:
                    run.font.size = Pt(7)


WORD_THEME = WordTheme(
    styles=STYLE_NAMES,
    setup=_prepare_document,
    layout=layout_hints,
    tables={"mfi_presentation": _add_presentation_table},
)
