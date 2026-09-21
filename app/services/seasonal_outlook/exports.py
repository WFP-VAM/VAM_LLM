"""Analyst exports with reproducible provenance and optional original map appendix."""
import io
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from .storage import encode


def word(state, maps, store, include_maps=False):
    document = Document()
    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    title = document.styles['Title']
    title.font.name = 'Arial'
    title.font.size = Pt(24)
    title.font.color.rgb = RGBColor(0, 0, 0)
    for border in title.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    section = document.sections[0]
    section.header.paragraphs[0].text = 'VAM LLM · Seasonal Outlook'
    report = state['report']
    document.add_heading(report['headline'], 0)
    document.add_paragraph(state['pack']['region_label'] + ' · ' + state['pack']['report_month'])
    for paragraph in report['paragraphs']:
        document.add_paragraph(paragraph['text'])
    if report['limitations']:
        document.add_heading('Limitations', 1)
        for item in report['limitations']:
            document.add_paragraph(item, style='List Bullet')
    document.add_heading('Evidence references', 1)
    for paragraph in report['paragraphs']:
        document.add_paragraph(paragraph['block_id'] + ': ' + ', '.join(paragraph['evidence_ids']))
    if include_maps:
        document.add_page_break()
        document.add_heading('Input maps', 1)
        for index, (figure, item) in enumerate(zip(state['pack']['figures'], maps), 1):
            if index > 1:
                document.add_page_break()
            document.add_heading(f'Figure {index} · {item["name"]}', 2)
            document.add_paragraph(figure['metadata_note'])
            data = store.read(item['object'])
            # Word cannot embed WebP; preserve the original in GCS/ZIP and
            # convert only its rendered appendix representation.
            with Image.open(io.BytesIO(data)) as image:
                width, height = image.size
                if image.format == 'WEBP' or width > 1860 or height > 1800:
                    # Bound Word's embedded raster to 300 dpi at its page size.
                    # This prevents twelve large WebP maps expanding beyond the
                    # worker's memory budget; originals remain in GCS and ZIP.
                    image.thumbnail((1860, 1800), Image.Resampling.LANCZOS)
                    converted = io.BytesIO()
                    rgba = image.convert('RGBA')
                    paper = Image.new('RGB', rgba.size, 'white')
                    paper.paste(rgba, mask=rgba.getchannel('A'))
                    paper.save(converted, format='PNG')
                    data = converted.getvalue()
            w = min(6.2, 6.0 * width / height)
            picture = document.add_picture(io.BytesIO(data), width=Inches(w))
            picture._inline.docPr.set('descr', figure['metadata_note'])
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def package(state, run, store, documents=None):
    result = io.BytesIO()
    with zipfile.ZipFile(result, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('input.json', encode(state['pack']))
        archive.writestr('provenance.json', encode({k: run[k] for k in ('id', 'region_id', 'report_date', 'confirmation', 'operations', 'versions')}))
        archive.writestr('rules.json', encode({k: state[k] for k in ('rules', 'evidence_profile', 'report_profile')}))
        for field in ('evidence_v1', 'review', 'evidence', 'issue_resolutions', 'analyst_comments',
                      'feedback_resolutions', 'initial_analysis', 'draft_review', 'report'):
            if field in state:
                archive.writestr(field + '.json', encode(state[field]))
        for version in run['versions']:
            archive.writestr('evidence_versions/' + version['id'] + '.json', store.read(version['object']))
        for operation in run['operations'].values():
            for i, ref in enumerate(operation['responses']):
                archive.writestr(f'attempts/{operation["id"]}/response-{i+1}.json', store.read(ref))
            for i, call in enumerate(operation['calls']):
                archive.writestr(f'attempts/{operation["id"]}/request-{i+1}.json', store.read(call['request']))
            for i, ref in enumerate(operation['checkpoints']):
                archive.writestr(f'attempts/{operation["id"]}/checkpoint-{i+1}.json', store.read(ref))
        for i, item in enumerate(run['maps'], 1):
            extension = {'image/png': '.png', 'image/jpeg': '.jpg', 'image/webp': '.webp'}[item['object']['mime']]
            archive.writestr(f'maps/{i:02}{extension}', store.read(item['object']))
        for name, data in (documents or {}).items():
            archive.writestr(name, data)
    return result.getvalue()


def build(state, run, store):
    documents = {'report.docx': word(state, run['maps'], store),
                 'report-with-maps.docx': word(state, run['maps'], store, True)}
    artifacts = {name: store.put(run['id'], data, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
                 for name, data in documents.items()}
    artifacts['artifacts.zip'] = store.put(run['id'], package(state, run, store, documents), 'application/zip')
    return artifacts
